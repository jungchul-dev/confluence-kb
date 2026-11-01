"""자료 작성 — 요구사항 추출 · 목차 제안 · 목차별 자료 수집 · 템플릿 적용 · 목차별 생성 · 요구사항 반영 점검.

흐름: 작성 건(이름 · 작성 개요 · 문서 형식 docx|pptx) → 요구사항 파일 → 항목 추출
     → 목차 제안(LLM) + 요구사항 배정 → 목차별 자료 자동 수집(기존 하이브리드 검색 + 리랭크) → 사람이 선별
     → (선택) 템플릿: 양식 파일(자리 표시 → 목차 · 내보내기 서식) / 참고 문서(소제목 → 목차, 생성 때 형식 예시)
       목차가 바뀌면 고른 자료 · 요구사항을 가장 가까운 새 목차로 옮긴다
     → 목차별 생성 → 내보내기(services/draft_export.py)
"""
import json
import logging
import math
import re
from dataclasses import dataclass
from typing import Dict, List, Optional, Tuple

from sqlalchemy import delete, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.db.models import Chunk, Document, Draft, DraftRequirement, DraftSection, DraftSource, Template
from app.permissions import CurrentUser
from app.services import answer as llm
from app.services.context import approx_tokens
from app.services.embeddings import embed_texts
from app.services.rerank import rerank
from app.services.retrieval import hybrid_search

logger = logging.getLogger(__name__)

AUTO_CANDIDATES = 10      # 목차마다 남기는 자동 후보 수
AUTO_SELECT = 4           # 그중 처음부터 선택해 두는 수
SECTION_CONTEXT_TOKENS = 9000  # 목차 하나를 생성할 때 넣는 자료 총량 상한
NO_THINK = {"chat_template_kwargs": {"enable_thinking": False}}
_PREFIX_LINE = re.compile(r"^\[[^\n]*\]\n?")


# ───────── 템플릿 목차 ─────────

@dataclass
class OutlineItem:
    title: str
    query: str = ""
    guide: str = ""
    placeholder: Optional[str] = None


_PLACEHOLDER = re.compile(r"\{\{\s*([^{}]+?)\s*\}\}")
RESERVED_PLACEHOLDERS = {"문서명", "작성일", "참고자료"}  # 목차가 아니라 내보낼 때 채우는 값 (services/draft_export.py)


def outline_from_file(file_name: str, data: bytes) -> List[OutlineItem]:
    """docx · pptx 템플릿 파일의 {{자리 표시}} 를 순서대로 목차로 만든다. 자리 표시가 없으면 docx 제목 스타일을 쓴다."""
    ext = file_name.rsplit(".", 1)[-1].lower() if "." in file_name else ""
    texts: List[Tuple[str, bool]] = []  # (텍스트, 제목 스타일 여부)
    import io
    if ext == "docx":
        from docx import Document as Docx  # type: ignore
        d = Docx(io.BytesIO(data))
        for p in _docx_all_paragraphs(d):
            style = (p.style.name or "") if p.style is not None else ""
            texts.append((p.text, bool(re.match(r"(Heading|제목) [1-3]$", style))))
    elif ext == "pptx":
        from pptx import Presentation  # type: ignore
        prs = Presentation(io.BytesIO(data))
        for slide in prs.slides:
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for para in shape.text_frame.paragraphs:
                        texts.append(("".join(r.text for r in para.runs), False))
    else:
        raise ValueError("템플릿 파일은 docx 또는 pptx 만 지원합니다.")

    items: List[OutlineItem] = []
    seen = set()
    for t, _ in texts:
        for m in _PLACEHOLDER.finditer(t):
            name = m.group(1).strip()
            if name and name not in seen and name not in RESERVED_PLACEHOLDERS:
                seen.add(name)
                items.append(OutlineItem(title=name, query=_strip_number(name), placeholder=name))
    if items:
        return items
    for t, is_heading in texts:  # 자리 표시가 없으면 제목 스타일 문단을 목차로 (내보낼 때는 새 문서로 만든다)
        t = t.strip()
        if is_heading and t and t not in seen:
            seen.add(t)
            items.append(OutlineItem(title=t, query=_strip_number(t)))
    if not items:
        raise ValueError("템플릿에서 {{자리 표시}} 나 제목 스타일 문단을 찾지 못했습니다.")
    return items


def _docx_all_paragraphs(d):
    for p in d.paragraphs:
        yield p
    for table in d.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    yield p


def _strip_number(title: str) -> str:
    return re.sub(r"^\s*([0-9IVX]+[.)]|[0-9]+(\.[0-9]+)*\.?|[가-하][.)])\s*", "", title).strip() or title


# ───────── 업로드 파일 파싱 ─────────

def _parse_sections(file_name: str, data: bytes):
    """업로드 파일 → 수집 파서(worker/parsers)의 Section 목록. 첨부 허용 확장자 설정과는 무관하게 동작한다."""
    from worker.parsers.office import parse_docx, parse_pptx, parse_text, parse_xlsx
    from worker.parsers.pdf import parse_pdf

    ext = file_name.rsplit(".", 1)[-1].lower() if "." in file_name else ""
    parsers = {"pdf": parse_pdf, "docx": parse_docx, "pptx": parse_pptx, "xlsx": parse_xlsx, "md": parse_text, "txt": parse_text, "csv": parse_text}
    fn = parsers.get(ext)
    if fn is None:
        raise ValueError(f"지원하지 않는 형식입니다: .{ext} (pdf · docx · pptx · xlsx · md · txt)")
    return fn(data)


def parse_upload(file_name: str, data: bytes) -> str:
    """요구사항 · 참고 문서 파일 → "## 소제목" 이 붙은 본문 텍스트."""
    parts = [(f"## {s.heading}\n" if s.heading else "") + s.text for s in _parse_sections(file_name, data)]
    return "\n\n".join(p for p in parts if p.strip())


_GENERIC_HEADING = re.compile(r"^(페이지 \d+|시트: .*|슬라이드 \d+)$")


def _clean_heading(h: str) -> str:
    return re.sub(r"^슬라이드 \d+:\s*", "", (h or "").strip()).strip()


def outline_from_reference(text: str) -> List[OutlineItem]:
    """참고 문서 본문의 "## 소제목" 을 순서대로 목차로 만든다. 페이지 · 시트 번호 같은 기계적 제목은 뺀다.
    쓸 만한 소제목이 2개 미만이면 빈 목록 (목차는 그대로 두고 형식 참고로만 쓴다)."""
    items: List[OutlineItem] = []
    seen = set()
    for line in (text or "").splitlines():
        if not line.startswith("## "):
            continue
        h = _clean_heading(line[3:])
        if not h or _GENERIC_HEADING.match(h) or h in seen:
            continue
        seen.add(h)
        items.append(OutlineItem(title=h[:300], query=_strip_number(h), guide="참고 문서의 같은 목차 형식 · 분량을 따른다"))
    return items if len(items) >= 2 else []


def _reference_parts(text: str) -> List[str]:
    parts, cur = [], []
    for line in (text or "").splitlines():
        if line.startswith("## ") and cur:
            parts.append("\n".join(cur).strip())
            cur = []
        cur.append(line)
    if cur:
        parts.append("\n".join(cur).strip())
    return [p for p in parts if p.strip()]


REFERENCE_EXCERPT_CHARS = 1000


async def reference_excerpt(ref_content: Optional[str], section: DraftSection) -> Optional[str]:
    """참고 문서에서 이 목차와 가장 가까운 부분 (형식 예시용). 제목이 같으면 그 부분, 아니면 임베딩으로 고른다."""
    parts = _reference_parts(ref_content or "")
    if not parts:
        return None
    key = _strip_number(section.title)
    for p in parts:
        first = p.splitlines()[0]
        if first.startswith("## ") and _strip_number(_clean_heading(first[3:])) == key:
            return p[:REFERENCE_EXCERPT_CHARS]
    if len(parts) == 1:
        return parts[0][:REFERENCE_EXCERPT_CHARS]
    try:
        vecs = await embed_texts([section_text(section)] + [p[:1500] for p in parts])
    except Exception as e:  # noqa: BLE001
        logger.warning("reference excerpt embedding failed: %s", e)
        return parts[0][:REFERENCE_EXCERPT_CHARS]
    best = max(range(1, len(vecs)), key=lambda i: _cos(vecs[0], vecs[i]))
    return parts[best - 1][:REFERENCE_EXCERPT_CHARS]


# ───────── 요구사항 추출 · 배정 ─────────

REQ_PROMPT = """당신은 요구사항 문서(요청서 · 과업지시서 · 작성 요청 메일 등)에서 '작성할 문서가 반드시 다뤄야 할 요구사항'을 뽑는 분석가입니다.
규칙:
1. 문서에 실제로 적힌 요구사항만 뽑습니다. 추측해서 만들지 않습니다.
2. 항목마다 원문 번호(예: SFR-001, 3.2.1)가 있으면 code 에, 없으면 null 로 둡니다.
3. content 는 원문 뜻을 유지한 한두 문장으로 씁니다. 수치·기한·조건은 그대로 옮깁니다.
4. 제출 방법 · 일정 안내처럼 문서 내용과 무관한 행정 안내는 빼되, 평가 · 검토 기준은 포함합니다.
5. JSON 배열만 출력합니다: [{"code": "SFR-001", "content": "..."}, ...]"""

_REQ_CHUNK_CHARS = 12000


async def extract_requirements(text: str) -> List[Dict[str, Optional[str]]]:
    """긴 문서는 나눠서 LLM 으로 항목을 뽑고 합친다. LLM 호출이 실패하면 번호 붙은 줄을 항목으로 쓴다."""
    text = text.strip()
    if not text:
        return []
    pieces = [text[i:i + _REQ_CHUNK_CHARS] for i in range(0, len(text), _REQ_CHUNK_CHARS)][:8]
    items: List[Dict[str, Optional[str]]] = []
    for piece in pieces:
        try:
            out = await llm.chat(REQ_PROMPT, f"문서:\n{piece}\n\n요구사항 JSON 배열:", temperature=0.0, max_tokens=4000, timeout=120.0, extra_body=NO_THINK)
            items += _parse_json_list(out)
        except Exception as e:  # noqa: BLE001
            logger.warning("requirement extraction via LLM failed, falling back to line split: %s", e)
            items += _fallback_requirements(piece)
    seen = set()
    uniq = []
    for it in items:
        key = (it.get("content") or "").strip()
        if key and key not in seen:
            seen.add(key)
            uniq.append({"code": (it.get("code") or None), "content": key})
    return uniq


def _parse_json_list(out: str) -> List[Dict[str, Optional[str]]]:
    m = re.search(r"\[.*\]", out or "", re.S)
    if not m:
        return []
    try:
        data = json.loads(m.group(0))
    except json.JSONDecodeError:
        return []
    res = []
    for d in data if isinstance(data, list) else []:
        if isinstance(d, dict) and str(d.get("content") or "").strip():
            code = d.get("code")
            res.append({"code": str(code).strip()[:64] if code else None, "content": str(d["content"]).strip()})
        elif isinstance(d, str) and d.strip():
            res.append({"code": None, "content": d.strip()})
    return res


def _fallback_requirements(text: str) -> List[Dict[str, Optional[str]]]:
    res = []
    for line in text.splitlines():
        m = re.match(r"^\s*((?:[A-Z]{2,5}-\d{2,4})|(?:\d+(?:\.\d+)+))\s*[.):]?\s+(.{8,})$", line)
        if m:
            res.append({"code": m.group(1), "content": m.group(2).strip()})
    return res


def _cos(a: List[float], b: List[float]) -> float:
    dot = sum(x * y for x, y in zip(a, b))
    na = math.sqrt(sum(x * x for x in a)) or 1.0
    nb = math.sqrt(sum(y * y for y in b)) or 1.0
    return dot / (na * nb)


def section_text(s: DraftSection) -> str:
    return "\n".join(t for t in (s.title, s.query, s.guide) if t)


async def assign_requirements(db: AsyncSession, draft: Draft, *, only_unassigned: bool = True) -> int:
    """요구사항 항목을 임베딩이 가장 가까운 목차에 배정한다. 임베딩이 없는 항목은 여기서 채운다."""
    sections = (await db.execute(select(DraftSection).where(DraftSection.draft_id == draft.id).order_by(DraftSection.ordinal))).scalars().all()
    reqs = (await db.execute(select(DraftRequirement).where(DraftRequirement.draft_id == draft.id))).scalars().all()
    if not sections or not reqs:
        return 0
    missing = [r for r in reqs if r.embedding is None]
    if missing:
        vecs = await embed_texts([r.content for r in missing])
        for r, v in zip(missing, vecs):
            r.embedding = v
    svecs = await embed_texts([section_text(s) for s in sections])
    n = 0
    for r in reqs:
        if only_unassigned and r.section_id is not None:
            continue
        rv = list(r.embedding)
        best = max(range(len(sections)), key=lambda i: _cos(rv, svecs[i]))
        if r.section_id != sections[best].id:
            # 다른 목차로 옮기면 예전 목차 본문 기준의 점검 결과는 맞지 않으므로 '점검 전' 으로 되돌린다
            r.covered, r.coverage_note = None, None
        r.section_id = sections[best].id
        n += 1
    await db.flush()
    return n


# ───────── 목차 제안 · 목차 바꾸기 ─────────

OUTLINE_PROMPT = """당신은 사내 문서의 목차를 짜는 편집자입니다. 문서 이름 · 작성 개요 · 요구사항 · (있으면) 참고 문서 목차를 보고 이 문서에 알맞은 목차를 제안합니다.
규칙:
1. 목차는 4~8개. 제목 앞에 "1." 처럼 번호를 붙입니다.
2. 요구사항이 있으면 모든 요구사항이 어느 한 목차에서 다뤄지도록 짭니다.
3. '참고 문서 목차' 가 있으면 그 구성 · 순서를 바탕으로 하되, 이번 문서의 개요 · 요구사항에 맞게 빼거나 더하고 이름을 고칩니다.
4. query 는 사내 문서에서 이 목차의 근거 자료를 찾을 검색어입니다 (명사 위주 3~8단어).
5. guide 는 이 목차에 쓸 내용 · 형식 · 분량 지침 한두 문장입니다.
6. JSON 배열만 출력합니다: [{"title": "1. 개요", "query": "...", "guide": "..."}]"""

FIXED_OUTLINE_PROMPT = """당신은 사내 문서의 목차를 다듬는 편집자입니다. 목차 제목은 템플릿으로 정해져 있어 바꿀 수 없습니다.
문서 이름 · 작성 개요 · 요구사항 · (있으면) 참고 문서 목차를 보고, 주어진 목차마다 찾을 내용과 작성 지침을 채웁니다.
규칙:
1. 주어진 목차를 같은 순서 · 같은 개수로, title 은 그대로 씁니다.
2. 요구사항이 있으면 각 요구사항을 가장 알맞은 목차의 guide 에 반영합니다.
3. query 는 사내 문서에서 이 목차의 근거 자료를 찾을 검색어입니다 (명사 위주 3~8단어).
4. guide 는 이 목차에 쓸 내용 · 형식 · 분량 지침 한두 문장입니다.
5. JSON 배열만 출력합니다: [{"title": "...", "query": "...", "guide": "..."}]"""

FALLBACK_OUTLINE = ["1. 개요", "2. 현황 및 배경", "3. 주요 내용", "4. 추진 계획", "5. 기대 효과"]


async def propose_outline(draft: Draft, reqs: List[DraftRequirement], *, fixed: Optional[List[OutlineItem]] = None,
                          reference: Optional[List[OutlineItem]] = None) -> List[OutlineItem]:
    """목차 제안. 요구사항 · 템플릿 · 참고 문서를 함께 본다.
    - fixed(템플릿의 자리 표시 목차)가 있으면 제목 · 순서 · 자리 표시는 그대로 두고 찾을 내용 · 작성 지침만 채운다 (내보낼 때 자리에 맞춰 넣어야 하므로)
    - reference(참고 문서 소제목)가 있으면 그 구성을 바탕으로 제안한다
    LLM 이 실패하면 fixed → reference → 일반 목차 순으로 그대로 쓴다."""
    parts = [f"문서 이름: {draft.name}"]
    if draft.overview.strip():
        parts.append(f"작성 개요:\n{draft.overview.strip()}")
    if reqs:
        parts.append("요구사항:\n" + "\n".join(f"- {('(' + r.code + ') ') if r.code else ''}{r.content}" for r in reqs[:60]))
    if reference:
        parts.append("참고 문서 목차:\n" + "\n".join(f"- {it.title}" for it in reference[:40]))
    if fixed:
        parts.append("정해진 목차:\n" + "\n".join(f"- {it.title}" for it in fixed))
    items: List[OutlineItem] = []
    try:
        out = await llm.chat(FIXED_OUTLINE_PROMPT if fixed else OUTLINE_PROMPT, "\n\n".join(parts) + "\n\n목차 JSON 배열:",
                             temperature=0.2, max_tokens=3000, timeout=90.0, extra_body=NO_THINK)
        m = re.search(r"\[.*\]", out or "", re.S)
        data = json.loads(m.group(0)) if m else []
        for d in data if isinstance(data, list) else []:
            if isinstance(d, dict) and str(d.get("title") or "").strip():
                t = str(d["title"]).strip()[:300]
                items.append(OutlineItem(title=t, query=str(d.get("query") or _strip_number(t)).strip()[:2000], guide=str(d.get("guide") or "").strip()[:4000]))
    except Exception as e:  # noqa: BLE001
        logger.warning("outline proposal failed, using fallback: %s", e)
    if fixed:
        # 제목 · 순서 · 자리 표시는 템플릿 그대로. LLM 결과는 같은 순서의 찾을 내용 · 지침만 가져온다
        by_title = {_strip_number(it.title): it for it in items}
        merged = []
        for i, f in enumerate(fixed):
            got = by_title.get(_strip_number(f.title)) or (items[i] if i < len(items) else None)
            merged.append(OutlineItem(title=f.title, query=(got.query if got else f.query) or f.query,
                                      guide=(got.guide if got else f.guide) or f.guide, placeholder=f.placeholder))
        return merged
    if items:
        return items[:12]
    if reference:
        return reference[:12]
    return [OutlineItem(title=t, query=f"{draft.name} {_strip_number(t)}") for t in FALLBACK_OUTLINE]


async def replace_outline(db: AsyncSession, draft: Draft, items: List[OutlineItem], template_section_ids: Optional[List] = None) -> int:
    """목차를 통째로 바꾼다. 고른 자료(selected)는 내용이 가장 가까운 새 목차로 옮기고, 고르지 않은 자동 후보와 생성 본문은 버린다.
    요구사항은 새 목차에 다시 배정한다. 옮긴 자료 수를 돌려준다."""
    old = (await db.execute(select(DraftSection).where(DraftSection.draft_id == draft.id))).scalars().all()
    new = [DraftSection(draft_id=draft.id, ordinal=i, title=it.title, query=it.query, guide=it.guide, space_keys=[], placeholder=it.placeholder,
                        template_section_id=(template_section_ids[i] if template_section_ids else None)) for i, it in enumerate(items)]
    db.add_all(new)
    await db.flush()

    keep = (await db.execute(select(DraftSource).where(DraftSource.draft_id == draft.id, DraftSource.selected.is_(True)))).scalars().all()
    moved = 0
    if keep and new:
        svecs = await embed_texts([section_text(s) for s in new])
        cvecs = await _source_vectors(db, keep)
        seen = set()
        for src, v in zip(keep, cvecs):
            best = max(range(len(new)), key=lambda i: _cos(v, svecs[i]))
            key = (best, src.document_id, src.chunk_ordinal)
            if key in seen:  # 여러 목차에서 같은 조각을 골랐으면 한 번만 남긴다
                await db.delete(src)
                continue
            seen.add(key)
            src.section_id = new[best].id
            moved += 1
    reqs = (await db.execute(select(DraftRequirement).where(DraftRequirement.draft_id == draft.id))).scalars().all()
    for r in reqs:
        r.section_id, r.covered, r.coverage_note = None, None, None
    await db.flush()
    old_ids = [s.id for s in old]
    if old_ids:
        await db.execute(delete(DraftSource).where(DraftSource.section_id.in_(old_ids)))
        await db.execute(delete(DraftSection).where(DraftSection.id.in_(old_ids)))
    await assign_requirements(db, draft, only_unassigned=False)
    return moved


async def _source_vectors(db: AsyncSession, sources: List[DraftSource]) -> List[List[float]]:
    """자료의 임베딩. 색인 때 저장해 둔 벡터(조각 = chunks.embedding, 문서 전체 = documents.embedding)를 그대로 쓰고,
    원문이 지워졌거나 재색인으로 조각 번호가 바뀌어 못 찾은 것만 새로 계산한다.
    (로컬 임베딩 서버는 긴 글을 한꺼번에 보내면 매우 느려서, 자료 26건을 새로 계산하면 80초 넘게 걸렸다)"""
    chunk_keys = {(s.document_id, s.chunk_ordinal) for s in sources if s.document_id and s.chunk_ordinal is not None}
    doc_ids = {s.document_id for s in sources if s.document_id and s.chunk_ordinal is None}
    by_chunk: Dict = {}
    if chunk_keys:
        rows = (await db.execute(
            select(Chunk.document_id, Chunk.ordinal, Chunk.embedding)
            .where(Chunk.document_id.in_({d for d, _ in chunk_keys}), Chunk.embedding.isnot(None))
        )).all()
        by_chunk = {(r.document_id, r.ordinal): list(r.embedding) for r in rows if (r.document_id, r.ordinal) in chunk_keys}
    by_doc: Dict = {}
    if doc_ids:
        rows = (await db.execute(select(Document.id, Document.embedding).where(Document.id.in_(doc_ids), Document.embedding.isnot(None)))).all()
        by_doc = {r.id: list(r.embedding) for r in rows}
    vecs: List[Optional[List[float]]] = [
        by_chunk.get((s.document_id, s.chunk_ordinal)) if s.chunk_ordinal is not None else by_doc.get(s.document_id) for s in sources
    ]
    missing = [i for i, v in enumerate(vecs) if v is None]
    if missing:
        fresh = await embed_texts([(sources[i].doc_title + "\n" + _PREFIX_LINE.sub("", sources[i].content).strip())[:1500] for i in missing])
        for i, v in zip(missing, fresh):
            vecs[i] = v
    return vecs  # type: ignore[return-value]


# ───────── 목차별 자료 수집 ─────────

def collect_query(draft: Draft, section: DraftSection, reqs: List[DraftRequirement]) -> str:
    """검색 질의 = 찾을 내용(없으면 제목) + 배정된 요구사항 일부 + 작성 개요 앞부분. 너무 길면 키워드 검색이 비므로 잘라 쓴다."""
    parts = [section.query.strip() or _strip_number(section.title)]
    if reqs:
        parts.append(" ".join(r.content for r in reqs[:3])[:300])
    if draft.overview.strip():
        parts.append(draft.overview.strip()[:150])
    return " ".join(parts)[:600]


async def collect_section(db: AsyncSession, draft: Draft, section: DraftSection, user: CurrentUser,
                          space_keys: Optional[List[str]] = None) -> int:
    """목차 하나의 자동 후보를 새로 모은다. 사람이 고르거나(selected) 직접 넣은(manual) 자료는 그대로 둔다.
    찾는 공간: 목차에 공간이 정해져 있으면 그것, 아니면 화면에서 고른 공간(space_keys), 둘 다 없으면 전체."""
    reqs = (await db.execute(select(DraftRequirement).where(DraftRequirement.section_id == section.id).order_by(DraftRequirement.ordinal))).scalars().all()
    q = collect_query(draft, section, list(reqs))
    candidates = await hybrid_search(db, q, user, space_keys=section.space_keys or space_keys or None)
    top, _ = await rerank(q, candidates, top_k=AUTO_CANDIDATES * 2)

    keep = (await db.execute(select(DraftSource).where(DraftSource.section_id == section.id))).scalars().all()
    kept_keys = {(s.document_id, s.chunk_ordinal) for s in keep if s.selected or s.origin != "auto"}
    await db.execute(delete(DraftSource).where(DraftSource.section_id == section.id, DraftSource.origin == "auto", DraftSource.selected.is_(False)))

    doc_hash = await _doc_hashes(db, {c.document_id for c in top})
    added = 0
    for c in top:
        if added >= AUTO_CANDIDATES:
            break
        if (c.document_id, c.ordinal) in kept_keys:
            continue
        db.add(DraftSource(
            draft_id=draft.id, section_id=section.id, origin="auto", document_id=c.document_id,
            doc_title=c.title, doc_path=c.path, doc_url=c.url, heading=c.heading, chunk_ordinal=c.ordinal,
            content=c.content, content_hash=doc_hash.get(c.document_id), score=c.rerank_score,
            selected=added < AUTO_SELECT and not any(s.selected for s in keep),
        ))
        added += 1
    await db.flush()
    return added


async def _doc_hashes(db: AsyncSession, ids) -> Dict:
    if not ids:
        return {}
    rows = (await db.execute(select(Document.id, Document.content_hash).where(Document.id.in_(list(ids))))).all()
    return {r.id: r.content_hash for r in rows}


async def add_manual_source(db: AsyncSession, draft: Draft, section: DraftSection, *, chunk_id=None, document_id=None, pinned: bool = False) -> DraftSource:
    """검색 결과의 조각(chunk_id) 또는 문서 전체(document_id) 를 내용 사본으로 넣는다."""
    if chunk_id:
        c = (await db.execute(select(Chunk).where(Chunk.id == chunk_id))).scalar_one_or_none()
        if not c:
            raise ValueError("조각을 찾을 수 없습니다. 재색인으로 바뀌었을 수 있으니 다시 검색하세요.")
        doc = (await db.execute(select(Document).where(Document.id == c.document_id))).scalar_one()
        content, heading, ordinal = c.content, c.heading, c.ordinal
    else:
        doc = (await db.execute(select(Document).where(Document.id == document_id))).scalar_one_or_none()
        if not doc:
            raise ValueError("문서를 찾을 수 없습니다.")
        chunks = (await db.execute(select(Chunk).where(Chunk.document_id == doc.id).order_by(Chunk.ordinal))).scalars().all()
        content = "\n\n".join(_PREFIX_LINE.sub("", x.content).strip() for x in chunks)
        heading, ordinal, pinned = None, None, True
    src = DraftSource(
        draft_id=draft.id, section_id=section.id, origin="manual", document_id=doc.id, doc_title=doc.title, doc_path=doc.path,
        doc_url=doc.url, heading=heading, chunk_ordinal=ordinal, content=content, content_hash=doc.content_hash,
        selected=True, pinned=pinned,
    )
    db.add(src)
    await db.flush()
    return src


def user_can_read(doc: Document, user: CurrentUser, skip: bool) -> bool:
    return skip or user.sees_everything or "*" in doc.allowed_groups or bool(set(doc.allowed_groups) & set(user.groups))


# ───────── 목차별 생성 ─────────

GEN_PROMPT = """당신은 사내 문서 작성자입니다. 주어진 '참고 자료'만 근거로 문서의 한 목차를 씁니다.
규칙:
1. 참고 자료에 있는 사실만 씁니다. 근거가 없는 내용(수치·일정·담당·제품명 등)은 지어내지 말고 [확인 필요: 무엇이 필요한지] 로 표시합니다.
2. 근거 문장 끝에 참고 자료 번호를 [1], [2] 형식으로 붙입니다.
3. 금액·일정·수량·SLA 같은 숫자는 원문 그대로 옮깁니다.
4. '요구사항' 이 주어지면 빠짐없이 다루고, 요구사항 번호가 있으면 본문에 (SFR-001) 처럼 밝힙니다.
5. 참고 자료에 나온 다른 고객 · 프로젝트 · 조직 이름을 이번 문서의 대상처럼 쓰지 않습니다. 이번 문서에 맞게 일반화하거나 작성 개요의 이름으로 바꿉니다.
6. 작성 지침의 형식·분량을 따릅니다. 표가 어울리면 마크다운 표를 씁니다.
7. 목차 제목은 다시 쓰지 말고 본문만 마크다운으로 출력합니다. 하위 제목이 필요하면 ### 를 씁니다.
8. '참고 문서 형식 예시' 는 다른 문서의 일부입니다. 구성 방식 · 문체 · 분량만 흉내 내고, 그 안의 문장 · 항목 · 이름 · 수치는 절대 옮기지 않습니다. 내용은 오직 '참고 자료' 와 '요구사항' 에서만 가져옵니다.
9. '문서 형식' 이 발표 자료(pptx)면 슬라이드에 들어갈 개조식 글머리표 위주로 짧게 씁니다. ### 소제목 하나가 슬라이드 한 장이 되며, 소제목마다 글머리표 3~6개, 항목은 한 줄 이내로 씁니다."""


def _source_block(i: int, s: DraftSource) -> str:
    body = _PREFIX_LINE.sub("", s.content).strip()
    where = s.doc_path + (f" > {s.heading}" if s.heading and s.chunk_ordinal is not None else "")
    return f"[{i}] {s.doc_title}\n경로: {where}\n---\n{body}"


async def generate_section(db: AsyncSession, draft: Draft, section: DraftSection) -> DraftSection:
    sources = (await db.execute(
        select(DraftSource).where(DraftSource.section_id == section.id, DraftSource.selected.is_(True))
        .order_by(DraftSource.pinned.desc(), DraftSource.score.desc().nulls_last(), DraftSource.created_at)
    )).scalars().all()
    reqs = (await db.execute(select(DraftRequirement).where(DraftRequirement.section_id == section.id).order_by(DraftRequirement.ordinal))).scalars().all()
    prev = (await db.execute(
        select(DraftSection).where(DraftSection.draft_id == draft.id, DraftSection.ordinal < section.ordinal, DraftSection.body.isnot(None))
        .order_by(DraftSection.ordinal)
    )).scalars().all()

    blocks, citations, used = [], [], 0
    for s in sources:
        block = _source_block(len(blocks) + 1, s)
        t = approx_tokens(block)
        if blocks and used + t > SECTION_CONTEXT_TOKENS:
            break
        if used + t > SECTION_CONTEXT_TOKENS:  # 첫 자료가 너무 길면 잘라서라도 넣는다
            block = block[:SECTION_CONTEXT_TOKENS]
            t = SECTION_CONTEXT_TOKENS
        blocks.append(block)
        used += t
        citations.append({"n": len(blocks), "source_id": str(s.id), "document_id": str(s.document_id) if s.document_id else None,
                          "title": s.doc_title, "path": s.doc_path, "heading": s.heading, "url": s.doc_url})

    parts = [f"문서: {draft.name}", "문서 형식: " + ("발표 자료 (pptx)" if draft.doc_format == "pptx" else "보고서 (docx)")]
    if draft.overview.strip():
        parts.append(f"작성 개요:\n{draft.overview.strip()}")
    if prev:
        parts.append("앞 목차 요약 (용어·수치를 맞출 것):\n" + "\n".join(f"- {p.title}: {_one_line(p.body)[:200]}" for p in prev[-6:]))
    parts.append(f"이번에 쓸 목차: {section.title}")
    if section.guide.strip():
        parts.append(f"작성 지침:\n{section.guide.strip()}")
    if reqs:
        parts.append("요구사항:\n" + "\n".join(f"- {('(' + r.code + ') ') if r.code else ''}{r.content}" for r in reqs))
    if draft.template_id:
        ref = (await db.execute(select(Template.ref_content).where(Template.id == draft.template_id))).scalar_one_or_none()
        excerpt = await reference_excerpt(ref, section) if ref else None
        if excerpt:
            parts.append("참고 문서 형식 예시 (다른 문서의 일부 — 형식만 참고하고 내용은 옮기지 말 것):\n<<<형식 예시 시작>>>\n"
                         f"{excerpt}\n<<<형식 예시 끝>>>")
    parts.append("참고 자료:\n\n" + ("\n\n=====\n\n".join(blocks) if blocks else "(선택된 자료 없음 — 모든 사실 내용은 [확인 필요] 로 표시)"))
    parts.append("본문:")

    body = await llm.chat(GEN_PROMPT, "\n\n".join(parts), temperature=0.2, max_tokens=4000, timeout=180.0, extra_body=NO_THINK)
    section.body = _strip_title(body, section.title)
    section.citations = citations
    section.version = (section.version or 0) + 1
    from datetime import datetime, timezone
    section.generated_at = datetime.now(timezone.utc)
    if reqs:
        await _check_coverage(section.body, list(reqs))
    await db.flush()
    return section


def _strip_title(body: str, title: str) -> str:
    """모델이 지시와 달리 첫 줄에 목차 제목을 다시 쓰면 지운다."""
    body = (body or "").strip()
    lines = body.splitlines()
    if lines:
        first = re.sub(r"^#{1,6}\s+", "", lines[0]).strip().strip("*").strip()
        if first and (_strip_number(first) == _strip_number(title) or first == title):
            body = "\n".join(lines[1:]).strip()
    return body


def _one_line(text: Optional[str]) -> str:
    return re.sub(r"\s+", " ", re.sub(r"\[\d+\]", "", text or "")).strip()


COVER_PROMPT = """본문이 각 요구사항을 실제로 다뤘는지 판정합니다. [확인 필요] 로만 표시된 것은 다루지 않은 것으로 봅니다.
요구사항의 수치·조건(예: 30분 이내, 99.5%)이 본문에 그대로 반영됐는지 확인합니다.
JSON 객체만 출력합니다: {"1": {"covered": true, "note": "한 문장 근거"}, "2": {"covered": false, "note": "빠진 점"}}"""


async def _check_coverage(body: str, reqs: List[DraftRequirement]) -> None:
    listing = "\n".join(f"{i}. {r.content}" for i, r in enumerate(reqs, start=1))
    try:
        out = await llm.chat(COVER_PROMPT, f"요구사항:\n{listing}\n\n본문:\n{body[:8000]}\n\n판정 JSON:", temperature=0.0, max_tokens=800, timeout=60.0, extra_body=NO_THINK)
        m = re.search(r"\{.*\}", out or "", re.S)
        verdict = _loose_json_object(m.group(0)) if m else {}
    except Exception as e:  # noqa: BLE001
        logger.warning("coverage check failed: %s", e)
        return
    for i, r in enumerate(reqs, start=1):
        v = verdict.get(str(i))
        if isinstance(v, dict):
            r.covered = bool(v.get("covered")) if v.get("covered") is not None else None
            r.coverage_note = str(v.get("note") or "")[:500] or None
        else:
            r.covered = bool(v) if v is not None else None
            r.coverage_note = None


def _loose_json_object(text: str) -> dict:
    """판정 JSON 뒤에 설명이 붙어 나와도 첫 번째 완전한 객체만 읽는다."""
    try:
        obj, _ = json.JSONDecoder().raw_decode(text)
        return obj if isinstance(obj, dict) else {}
    except json.JSONDecodeError:
        return {}
