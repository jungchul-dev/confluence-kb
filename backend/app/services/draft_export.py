"""작성 건 → docx / pptx (작성 건의 문서 형식을 따른다).

- 템플릿 파일이 있으면 그 파일의 {{자리 표시}} 를 목차 본문으로 채운다 (서식·표지는 템플릿 그대로).
  예약 자리 표시: {{문서명}} {{작성일}} {{참고자료}}
- 양식 파일이 없으면 새로 만든다. docx: 제목 → 목차별 제목·본문 → 참고 자료, pptx: 표지 → 목차별 슬라이드 → 참고 자료.
- 목차마다 [1] [2] 로 매긴 출처 번호를 문서 전체 기준 번호로 다시 매기고, 끝에 참고 자료 목록을 붙인다.
"""
import io
import re
from datetime import date
from typing import Dict, List, Optional, Tuple

from app.db.models import Draft, DraftSection

RESERVED = {"문서명", "작성일", "참고자료"}
_CITE = re.compile(r"\[(\d+)\]")


def renumber(sections: List[DraftSection]) -> Tuple[Dict[str, str], List[dict]]:
    """목차별 출처 번호 → 전체 번호. (목차 id → 번호를 바꾼 본문, 참고 자료 목록) 을 돌려준다."""
    refs: List[dict] = []
    key_to_n: Dict[str, int] = {}
    bodies: Dict[str, str] = {}
    for s in sections:
        local: Dict[int, int] = {}
        for c in s.citations or []:
            key = c.get("document_id") or c.get("url") or c.get("title")
            if key not in key_to_n:
                refs.append(c)
                key_to_n[key] = len(refs)
            local[int(c["n"])] = key_to_n[key]
        body = _CITE.sub(lambda m: f"[{local.get(int(m.group(1)), m.group(1))}]", s.body or "")
        bodies[str(s.id)] = re.sub(r"(\[\d+\])(?:\1)+", r"\1", body)  # 같은 문서라 [1][1] 로 겹치면 하나로
    return bodies, refs


def ref_line(n: int, r: dict) -> str:
    return f"[{n}] {r.get('title') or ''} — {r.get('path') or ''}" + (f" ({r['url']})" if r.get("url") else "")


# ───────── 마크다운 → docx 블록 ─────────

def _blocks(md: str) -> List[Tuple[str, object]]:
    """('h', text) · ('p', text) · ('ul', text) · ('ol', text) · ('table', rows) 목록."""
    out: List[Tuple[str, object]] = []
    lines = (md or "").splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line.strip():
            i += 1
            continue
        if line.lstrip().startswith("|"):
            rows = []
            while i < len(lines) and lines[i].lstrip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
                    rows.append(cells)
                i += 1
            if rows:
                out.append(("table", rows))
            continue
        m = re.match(r"^(#{1,6})\s+(.*)$", line)
        if m:
            out.append(("h", m.group(2).strip()))
        elif re.match(r"^\s*[-*•]\s+", line):
            out.append(("ul", re.sub(r"^\s*[-*•]\s+", "", line)))
        elif re.match(r"^\s*\d+[.)]\s+", line):
            out.append(("ol", re.sub(r"^\s*\d+[.)]\s+", "", line)))
        else:
            out.append(("p", line.strip()))
        i += 1
    return out


def _add_runs(p, text: str) -> None:
    for part in re.split(r"(\*\*[^*]+\*\*)", text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            p.add_run(part[2:-2]).bold = True
        else:
            p.add_run(part)


def _style_or_none(doc, name: str):
    try:
        return doc.styles[name]
    except KeyError:
        return None


def _make_elements(doc, md: str) -> list:
    """마크다운 본문을 docx 문단·표 요소로 만든다. 문서 끝에 붙은 요소를 돌려주므로 호출 쪽에서 원하는 위치로 옮긴다."""
    els = []
    for kind, val in _blocks(md):
        if kind == "table":
            rows = val
            ncol = max(len(r) for r in rows)
            t = doc.add_table(rows=len(rows), cols=ncol)
            if _style_or_none(doc, "Table Grid") is not None:
                t.style = doc.styles["Table Grid"]
            for ri, r in enumerate(rows):
                for ci in range(ncol):
                    cell = t.cell(ri, ci)
                    cell.text = ""
                    _add_runs(cell.paragraphs[0], r[ci] if ci < len(r) else "")
                    if ri == 0:
                        for run in cell.paragraphs[0].runs:
                            run.bold = True
            els.append(t._tbl)
            continue
        style = {"h": "Heading 3", "ul": "List Bullet", "ol": "List Number"}.get(kind)
        p = doc.add_paragraph()
        if style and _style_or_none(doc, style) is not None:
            p.style = doc.styles[style]
        elif kind == "ul":
            val = "• " + str(val)
        _add_runs(p, str(val))
        if kind == "h" and p.style.name != "Heading 3":
            for run in p.runs:
                run.bold = True
        els.append(p._p)
    return els


def _docx_paragraphs(doc):
    for p in doc.paragraphs:
        yield p
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    yield p


def export_docx(draft: Draft, sections: List[DraftSection], template_bytes: Optional[bytes]) -> bytes:
    from docx import Document as Docx  # type: ignore

    bodies, refs = renumber(sections)
    refs_md = "\n".join(f"- {ref_line(i, r)}" for i, r in enumerate(refs, start=1))
    by_placeholder = {(s.placeholder or s.title): bodies[str(s.id)] for s in sections}
    simple = {"문서명": draft.name, "작성일": date.today().isoformat()}

    if template_bytes:
        doc = Docx(io.BytesIO(template_bytes))
        for p in list(_docx_paragraphs(doc)):
            text = p.text
            names = [m.strip() for m in re.findall(r"\{\{\s*([^{}]+?)\s*\}\}", text)]
            if not names:
                continue
            # 문단 전체가 자리 표시 하나면 본문 블록(문단·표·목록)으로 통째로 바꾼다
            if len(names) == 1 and re.fullmatch(r"\s*\{\{\s*[^{}]+?\s*\}\}\s*", text) and names[0] not in simple:
                name = names[0]
                md = refs_md if name == "참고자료" else by_placeholder.get(name)
                if md is None:
                    continue
                for el in _make_elements(doc, md or "[확인 필요: 생성되지 않음]"):
                    p._p.addprevious(el)
                p._p.getparent().remove(p._p)
                continue
            # 문장 중간의 자리 표시는 글자로만 바꾼다 (서식은 첫 run 을 따른다)
            new = re.sub(r"\{\{\s*([^{}]+?)\s*\}\}", lambda m: simple.get(m.group(1).strip(), _plain(by_placeholder.get(m.group(1).strip(), m.group(0)))), text)
            for r in p.runs[1:]:
                r.text = ""
            if p.runs:
                p.runs[0].text = new
        buf = io.BytesIO()
        doc.save(buf)
        return buf.getvalue()

    doc = Docx()
    doc.add_heading(draft.name, level=0)
    doc.add_paragraph(f"작성일 {date.today().isoformat()}")
    if draft.overview.strip():
        doc.add_heading("작성 개요", level=1)
        for line in draft.overview.strip().splitlines():
            if line.strip():
                doc.add_paragraph(line.strip())
    for s in sections:
        doc.add_heading(s.title, level=1)
        _make_elements(doc, bodies[str(s.id)] or "[확인 필요: 생성되지 않음]")
    if refs:
        doc.add_heading("참고 자료", level=1)
        for i, r in enumerate(refs, start=1):
            doc.add_paragraph(ref_line(i, r))
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _plain(md: str) -> str:
    text = re.sub(r"^#{1,6}\s+", "", md or "", flags=re.M)
    text = re.sub(r"\*\*([^*]+)\*\*", r"\1", text)
    return text.strip()


def export_pptx(draft: Draft, sections: List[DraftSection], template_bytes: bytes) -> bytes:
    """pptx 템플릿의 {{자리 표시}} 가 든 텍스트 상자를 본문(글자만)으로 채운다. 표는 '열1 | 열2' 줄로 들어간다."""
    from pptx import Presentation  # type: ignore

    bodies, refs = renumber(sections)
    by_placeholder = {(s.placeholder or s.title): _plain(bodies[str(s.id)]) for s in sections}
    values = {"문서명": draft.name, "작성일": date.today().isoformat(),
              "참고자료": "\n".join(ref_line(i, r) for i, r in enumerate(refs, start=1)), **by_placeholder}
    prs = Presentation(io.BytesIO(template_bytes))
    for slide in prs.slides:
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            tf = shape.text_frame
            full = "\n".join("".join(r.text for r in p.runs) for p in tf.paragraphs)
            if "{{" not in full:
                continue
            new = re.sub(r"\{\{\s*([^{}]+?)\s*\}\}", lambda m: values.get(m.group(1).strip(), m.group(0)), full)
            lines = [ln.strip("|").strip() if ln.strip().startswith("|") else ln for ln in new.splitlines()]
            lines = [ln for ln in lines if not re.fullmatch(r"[\s|:-]*-{2,}[\s|:-]*", ln)]
            first = tf.paragraphs[0]
            font_runs = first.runs[:1]
            for p in tf.paragraphs[1:]:
                p._p.getparent().remove(p._p)
            for r in first.runs[1:]:
                r._r.getparent().remove(r._r)
            if font_runs:
                font_runs[0].text = lines[0] if lines else ""
            else:
                first.text = lines[0] if lines else ""
            for ln in lines[1:]:
                para = tf.add_paragraph()
                para.text = ln
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


# ───────── 양식 없이 새 pptx ─────────
# 표지 → 작성 개요 → 목차마다 슬라이드(### 소제목마다 한 장, 글머리표가 많으면 나눔 · 표는 따로 한 장) → 참고 자료

_BULLETS_PER_SLIDE = 8


def _pptx_runs(paragraph, text: str, size) -> None:
    for part in re.split(r"(\*\*[^*]+\*\*)", text):
        if not part:
            continue
        run = paragraph.add_run()
        bold = part.startswith("**") and part.endswith("**")
        run.text = part[2:-2] if bold else part
        run.font.size = size
        run.font.bold = bold or None


def export_pptx_new(draft: Draft, sections: List[DraftSection]) -> bytes:
    from pptx import Presentation  # type: ignore
    from pptx.util import Inches, Pt  # type: ignore

    bodies, refs = renumber(sections)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    title_layout, content_layout, title_only = prs.slide_layouts[0], prs.slide_layouts[1], prs.slide_layouts[5]

    def fit(shape, top=Inches(1.5)):
        shape.left, shape.top, shape.width, shape.height = Inches(0.6), top, prs.slide_width - Inches(1.2), prs.slide_height - top - Inches(0.5)

    def bullet_slide(title: str, lines: List[str]) -> None:
        slide = prs.slides.add_slide(content_layout)
        slide.shapes.title.text = title
        slide.shapes.title.text_frame.paragraphs[0].runs[0].font.size = Pt(30)
        fit(slide.shapes.title, Inches(0.3))
        slide.shapes.title.height = Inches(1.0)
        body = slide.placeholders[1]
        fit(body)
        tf = body.text_frame
        tf.word_wrap = True
        for i, ln in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            _pptx_runs(p, ln, Pt(18))

    def table_slide(title: str, rows: List[List[str]]) -> None:
        slide = prs.slides.add_slide(title_only)
        slide.shapes.title.text = title
        slide.shapes.title.text_frame.paragraphs[0].runs[0].font.size = Pt(30)
        fit(slide.shapes.title, Inches(0.3))
        slide.shapes.title.height = Inches(1.0)
        ncol = max(len(r) for r in rows)
        shape = slide.shapes.add_table(len(rows), ncol, Inches(0.6), Inches(1.5), prs.slide_width - Inches(1.2), Inches(0.4) * len(rows))
        for ri, r in enumerate(rows):
            for ci in range(ncol):
                cell = shape.table.cell(ri, ci)
                cell.text = ""
                _pptx_runs(cell.text_frame.paragraphs[0], re.sub(r"\*\*([^*]+)\*\*", r"\1", r[ci] if ci < len(r) else ""), Pt(14))

    def emit(title: str, lines: List[str]) -> None:
        for i in range(0, max(len(lines), 1), _BULLETS_PER_SLIDE):
            bullet_slide(title + (" (계속)" if i else ""), lines[i:i + _BULLETS_PER_SLIDE] or ["[확인 필요: 생성되지 않음]"])

    cover = prs.slides.add_slide(title_layout)
    cover.shapes.title.text = draft.name
    cover.placeholders[1].text = f"작성일 {date.today().isoformat()}"
    for shp in (cover.shapes.title, cover.placeholders[1]):
        shp.left, shp.width = Inches(0.6), prs.slide_width - Inches(1.2)

    overview = [ln.strip() for ln in (draft.overview or "").splitlines() if ln.strip()]
    if overview:
        emit("작성 개요", overview)

    for s in sections:
        title, lines, emitted = s.title, [], False
        for kind, val in _blocks(bodies[str(s.id)]):
            if kind == "h":
                if lines:
                    emit(title, lines)
                    emitted = True
                title, lines = f"{s.title} · {val}", []
            elif kind == "table":
                if lines:
                    emit(title, lines)
                    lines = []
                table_slide(title, val)
                emitted = True
            else:
                lines.append(str(val))
        if lines or not emitted:
            emit(title, lines)

    if refs:
        emit("참고 자료", [ref_line(i, r) for i, r in enumerate(refs, start=1)])
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()
